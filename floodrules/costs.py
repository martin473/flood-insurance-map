"""What neighbors paid: premium summaries built from published federal datasets.

No model and no scraping logic, just two fixed downloads:

* Flood: FEMA's NFIP Redacted Policies v3 (OpenFEMA API), every federal flood policy with its
  cost, flood zone and census block group. We keep single-family, non-condo, one-year policies
  that started in the most recent 12 months of data, and summarize what they cost
  (premium plus the fees and surcharges on the bill) by census tract and flood zone, with town
  and county as fallbacks for tracts with too few policies.
* Homeowners: the Treasury Federal Insurance Office's ZIP-level homeowners data (2018-2022),
  removed from Treasury's site in 2025 and re-hosted by the Consumer Federation of America.
  It publishes an average premium per policy per ZIP, so that's what we report.

The output is one small JSON file the map loads: map/data/ri-costs.json.
"""

from __future__ import annotations

import io
import json
import re
import statistics
from collections import defaultdict
from collections.abc import Iterable, Iterator
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import requests

FEMA_POLICIES = "https://www.fema.gov/api/open/v3/NfipPolicies"
FEMA_DATASET_PAGE = "https://www.fema.gov/openfema-data-page/nfip-redacted-policies-v3"
FIO_XLSX = (
    "https://consumerfed.org/media/legacy/post_31852/"
    "Supporting_Underlying_Metrics_and_Disclaimer_for_Analyses_of_US_Homeowners_Insurance_Markets_2018-2022-2.xlsx"
)
FIO_PAGE = (
    "https://consumerfed.org/news/blogs/download-analyses-of-us-homeowners-insurance-markets-2018-2022-"
    "climate-related-risks-and-other-factors-report-and-data-from-federal-insurance-office-fio/"
)

FIELDS = [
    "policyEffectiveDate", "policyCost", "totalInsurancePremiumOfThePolicy", "ratedFloodZone",
    "censusGeoid", "occupancyType", "condominiumCoverageTypeCode", "policyTermIndicator", "nfipCommunityName",
]
SINGLE_FAMILY = {1, 11}       # 1 = legacy rating, 11 = Risk Rating 2.0
MIN_POLICIES = 10             # fewer than this and a group isn't shown
PAGE = 10000                  # OpenFEMA's maximum page size


# ---------- download ----------

def fetch_policies(state: str, since: date, session: requests.Session | None = None, page: int = PAGE) -> Iterator[dict]:
    """Every policy record for a state that started on or after `since`, paged from OpenFEMA."""
    s = session or requests.Session()
    flt = f"propertyState eq '{state}' and policyEffectiveDate ge '{since.isoformat()}'"
    skip = 0
    while True:
        params: dict[str, str | int] = {"$filter": flt, "$select": ",".join(FIELDS), "$top": page, "$skip": skip, "$orderby": "id"}
        r = s.get(FEMA_POLICIES, params=params, timeout=120)
        r.raise_for_status()
        rows = r.json().get("NfipPolicies", [])
        yield from rows
        if len(rows) < page:
            return
        skip += page


def fetch_homeowners_rows(session: requests.Session | None = None) -> list[tuple]:
    """Rows of the FIO 'Supporting Underlying Metrics' sheet: (ZIP, year, decile, ..., premium per policy, nonrenewal rate, ...)."""
    from openpyxl import load_workbook  # optional dependency: pip install -e .[data]

    r = (session or requests.Session()).get(FIO_XLSX, timeout=180)
    r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["Supporting Underlying Metrics"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip() for h in rows[0]]
    want = ["ZIP Code", "Year", "Premiums Per Policy", "Nonrenewal Rate"]
    missing = [w for w in want if w not in header]
    if missing:
        raise ValueError(f"FIO spreadsheet columns changed, missing {missing}; found {header}")
    idx = [header.index(w) for w in want]
    return [tuple(row[i] for i in idx) for row in rows[1:]]


# ---------- summarize ----------

def zone_group(zone: str | None) -> str | None:
    """'high' for FEMA's high-risk zones (A and V, where lenders require flood insurance), 'other' for the rest."""
    z = (zone or "").strip().upper()
    if not z:
        return None
    if z.startswith(("A", "V")) and z not in ("A99", "AR"):
        return "high"
    return "other"


def town_name(community: str | None) -> str | None:
    """'SOUTH KINGSTOWN, TOWN OF' -> 'south kingstown'. Matches the town names the map gets from OpenStreetMap."""
    if not community:
        return None
    name = re.sub(r",\s*(TOWN|CITY) OF\s*$", "", community.strip(), flags=re.I)
    return name.lower() or None


def _day(s: str | None) -> date | None:
    try:
        return datetime.fromisoformat((s or "").replace("Z", "+00:00")).date()
    except ValueError:
        return None


def keep(p: dict) -> bool:
    """Single-family homes, not condos, one-year policies with a positive bill."""
    return (p.get("occupancyType") in SINGLE_FAMILY
            and (p.get("condominiumCoverageTypeCode") or "N") == "N"
            and p.get("policyTermIndicator") == 1
            and (p.get("policyCost") or 0) > 0
            and (p.get("totalInsurancePremiumOfThePolicy") or 0) > 0)


def stats(values: list[float]) -> dict[str, int]:
    v = sorted(values)
    q = statistics.quantiles(v, n=4) if len(v) >= 2 else [v[0], v[0], v[0]]
    return {"n": len(v), "median": round(statistics.median(v)), "mean": round(statistics.fmean(v)),
            "p25": round(q[0]), "p75": round(q[2])}


def summarize_flood(policies: Iterable[dict], months: int = 12, min_n: int = MIN_POLICIES, today: date | None = None) -> dict[str, Any]:
    """One year of policies, so each insured home counts about once (they renew yearly).

    The window ends at today, or at the newest policy if FEMA's data stops earlier. Renewals
    booked ahead with future start dates are left out."""
    today = today or datetime.now(timezone.utc).date()
    rows = []
    for p in policies:
        d = _day(p.get("policyEffectiveDate"))
        g = zone_group(p.get("ratedFloodZone"))
        geo = str(p.get("censusGeoid") or "")
        if d and g and keep(p):
            rows.append((d, g, geo, town_name(p.get("nfipCommunityName")), float(p["policyCost"])))
    if not rows:
        raise ValueError("no usable policies; the dataset or its fields may have changed")
    rows = [r for r in rows if r[0] <= today]
    if not rows:
        raise ValueError("no policies on or before today")
    latest = max(r[0] for r in rows)
    start = latest - timedelta(days=round(months * 365.25 / 12)) + timedelta(days=1)
    rows = [r for r in rows if r[0] >= start]

    groups: dict[str, dict[str, list[float]]] = {"tract": defaultdict(list), "town": defaultdict(list), "county": defaultdict(list)}
    for _, g, geo, town, cost in rows:
        if len(geo) >= 11:
            groups["tract"][f"{geo[:11]}|{g}"].append(cost)
            groups["county"][f"{geo[:5]}|{g}"].append(cost)
        if town:
            groups["town"][f"{town}|{g}"].append(cost)
    out: dict[str, Any] = {
        "window": {"from": start.isoformat(), "to": latest.isoformat()},
        "policies": len(rows),
        "min_policies": min_n,
    }
    for level, by in groups.items():
        out[level] = {k: stats(v) for k, v in sorted(by.items()) if len(v) >= min_n}
    out["statewide"] = {g: stats([r[4] for r in rows if r[1] == g]) for g in ("high", "other") if any(r[1] == g for r in rows)}
    return out


def summarize_homeowners(rows: Iterable[tuple], zip_lo: int, zip_hi: int) -> dict[str, Any]:
    """Latest year's average premium per policy for each ZIP in [zip_lo, zip_hi]. ZIPs come as numbers (leading zero dropped)."""
    by_zip: dict[str, tuple[int, float, float | None]] = {}
    for z, year, prem, nonrenew in rows:
        try:
            zi, yi, pf = int(z), int(year), float(prem)
        except (TypeError, ValueError):
            continue
        if not (zip_lo <= zi <= zip_hi) or pf <= 0:
            continue
        key = f"{zi:05d}"
        if key not in by_zip or yi > by_zip[key][0]:
            by_zip[key] = (yi, pf, float(nonrenew) if nonrenew is not None else None)
    if not by_zip:
        raise ValueError("no ZIPs in range; the spreadsheet layout may have changed")
    return {
        "year": max(v[0] for v in by_zip.values()),
        "zip": {k: {"year": y, "avg": round(p), "nonrenewal_pct": None if n is None else round(n * 100, 2)}
                for k, (y, p, n) in sorted(by_zip.items())},
    }


def build(state: str = "RI", zip_range: tuple[int, int] = (2801, 2940), today: date | None = None,
          session: requests.Session | None = None) -> dict[str, Any]:
    today = today or datetime.now(timezone.utc).date()
    # two years back, so the newest 12 months are covered even with FEMA's reporting lag
    flood = summarize_flood(fetch_policies(state, today - timedelta(days=730), session), today=today)
    homeowners = summarize_homeowners(fetch_homeowners_rows(session), *zip_range)
    return {
        "built_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "state": state,
        "flood": {**flood, "source": "FEMA NFIP Redacted Policies v3 (OpenFEMA)", "url": FEMA_DATASET_PAGE,
                  "what": "Total yearly cost (premium plus federal fees and surcharges) of federal flood policies on single-family homes."},
        "homeowners": {**homeowners, "source": "U.S. Treasury Federal Insurance Office, homeowners data 2018-2022 (via Consumer Federation of America)",
                       "url": FIO_PAGE, "what": "Average homeowners (HO-3/HO-5) premium per policy from insurers that reported; excludes the FAIR Plan."},
    }


def check(data: dict[str, Any], min_policies: int = 2000, min_tracts: int = 50, min_zips: int = 50) -> list[str]:
    """Problems that mean the download went wrong (a truncated page, a renamed field). Empty when the data looks sane."""
    f, h = data["flood"], data["homeowners"]
    problems = []
    if f["policies"] < min_policies:
        problems.append(f"only {f['policies']} flood policies (expected at least {min_policies})")
    if len(f["tract"]) < min_tracts:
        problems.append(f"only {len(f['tract'])} census tract groups (expected at least {min_tracts})")
    if len(h["zip"]) < min_zips:
        problems.append(f"only {len(h['zip'])} homeowners ZIPs (expected at least {min_zips})")
    for g, st in f["statewide"].items():
        if not 100 <= st["median"] <= 20000:
            problems.append(f"statewide {g} median ${st['median']} is implausible")
    return problems


def write(data: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, separators=(",", ":"), sort_keys=True) + "\n")
